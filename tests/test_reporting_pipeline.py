from __future__ import annotations

from io import BytesIO
import json
import zipfile

import numpy as np
import pandas as pd
import pytest
from docx import Document

from analysis import (
    categorical_summary, group_statistics_detailed, numeric_summary,
    paired_change_statistics, pearson_relationship,
)
from inspector.charts import render_chart
from inspector.exports import export_docx, export_html
from inspector.ingest import ParsedData, infer_roles
from inspector.pipeline import ReportSettings, generate_report


def parsed_fixture() -> ParsedData:
    data = pd.DataFrame({
        "participant_id": range(1, 13),
        "group": ["A"] * 6 + ["B"] * 6,
        "score": [1.0, 2.0, 3.0, 4.0, np.inf, np.nan, 4.0, 5.0, 6.0, 7.0, 8.0, 9.0],
        "post": [2.0, 3.0, 4.0, 5.0, 6.0, np.nan, 5.0, 6.0, 7.0, 8.0, 9.0, 10.0],
        "date": pd.date_range("2026-01-01", periods=12),
    })
    ids = {name: f"c{i:03d}" for i, name in enumerate(data.columns, 1)}
    return ParsedData(data=data, notes=["synthetic"], column_ids=ids)


def test_numeric_summary_handles_infinity_and_singletons() -> None:
    result = numeric_summary(pd.DataFrame({"x": [1.0, np.inf, np.nan]}), "x")
    assert result["n"] == 1
    assert result["sd"] is None
    assert result["missing_count"] == 1
    assert result["nonfinite_count"] == 1


def test_categorical_percentages_keep_full_denominator_and_other() -> None:
    data = pd.DataFrame({"x": [f"c{i}" for i in range(22)] + [None]})
    result = categorical_summary(data, "x", max_categories=20)
    assert result["valid_n"] == 22
    assert result["categories"][-1]["label"] == "Other"
    assert sum(row["percent"] for row in result["categories"]) == pytest.approx(100)


def test_group_complete_cases_keep_group_and_direction_gate() -> None:
    parsed = parsed_fixture()
    result = group_statistics_detailed(parsed.data, "score", "group")
    assert result["valid_n"] == 10
    assert result["group_order"] == ["A", "B"]
    output = generate_report(parsed, ReportSettings(report_type="group", outcome="score", group="group", role_overrides={"score": "numeric"}))
    assert "effect" not in output.analysis.selected_analyses["group"]
    assert "independent" in output.analysis.selected_analyses["group"]["effect_reason"]


def test_pearson_pairwise_and_constant_reason() -> None:
    data = pd.DataFrame({"x": [1, 2, 3, np.nan], "y": [2, 4, 6, 8]})
    result = pearson_relationship(data, "x", "y")
    assert result["n"] == 3 and result["r"] == pytest.approx(1)
    constant = pearson_relationship(pd.DataFrame({"x": [1, 1, 1], "y": [1, 2, 3]}), "x", "y")
    assert constant["r"] is None and constant["reason"] == "zero_variance"


def test_paired_change_uses_identical_subset_and_post_minus_pre() -> None:
    data = pd.DataFrame({"pre": [1.0, 2.0, np.nan], "post": [3.0, np.nan, 8.0]})
    result = paired_change_statistics(data, "pre", "post")
    assert result["paired_n"] == 1
    assert result["change_summary"]["mean"] == pytest.approx(2)
    assert result["direction"] == "post_minus_pre"


@pytest.mark.parametrize("language", ["en", "zh-CN", "es", "fi"])
def test_default_pipeline_is_localized_json_safe_and_renders_charts(language: str) -> None:
    parsed = parsed_fixture()
    output = generate_report(parsed, ReportSettings(language=language, role_overrides={"score": "numeric", "post": "numeric"}))
    assert output.report.language == language
    assert output.report.mode == "template"
    assert output.report.sections
    serialized = output.analysis.model_dump_json()
    assert "NaN" not in serialized and "Infinity" not in serialized
    for section in output.report.sections:
        for figure in section.figures:
            assert figure.png.startswith(b"\x89PNG")


def test_chart_budget_and_scatter_sampling_disclosure() -> None:
    parsed = parsed_fixture()
    output = generate_report(parsed, ReportSettings(report_type="relationships", relationship_first="score", relationship_second="post", role_overrides={"score": "numeric", "post": "numeric"}, depth="brief"))
    assert len(output.analysis.chart_specs) <= 3
    primary = output.analysis.chart_specs[0]
    assert primary.kind == "scatter" and "5,000" in (primary.sampling or "")


def test_group_report_requires_selectors_and_change_requires_confirmation() -> None:
    parsed = parsed_fixture()
    with pytest.raises(Exception, match="required"):
        generate_report(parsed, ReportSettings(report_type="group"))
    with pytest.raises(Exception, match="Confirm"):
        generate_report(parsed, ReportSettings(report_type="change", pre="score", post="post", role_overrides={"score": "numeric", "post": "numeric"}))


def test_docx_and_html_share_text_tables_and_embedded_images() -> None:
    output = generate_report(parsed_fixture(), ReportSettings(language="zh-CN", role_overrides={"score": "numeric", "post": "numeric"}))
    docx_bytes = export_docx(output.report)
    document = Document(BytesIO(docx_bytes))
    assert output.report.title in "\n".join(paragraph.text for paragraph in document.paragraphs)
    assert document.tables
    with zipfile.ZipFile(BytesIO(docx_bytes)) as archive:
        assert any(name.startswith("word/media/") for name in archive.namelist())
        assert "word/fonts/NotoSansCJKsc-Regular.odttf" in archive.namelist()
        assert b"embedRegular" in archive.read("word/fontTable.xml")
    html_bytes = export_html(output.report)
    text = html_bytes.decode("utf-8")
    assert "@media print" in text and "data:image/png;base64" in text
    assert output.report.title in text


def test_html_escapes_hostile_labels() -> None:
    parsed = ParsedData(data=pd.DataFrame({"<script>alert(1)</script>": [1, 2, 3, 4, 5]}), notes=[], column_ids={"<script>alert(1)</script>": "c1"})
    report = generate_report(parsed, ReportSettings(role_overrides={"<script>alert(1)</script>": "numeric"})).report
    rendered = export_html(report).decode()
    assert "<script>alert(1)</script>" not in rendered
    assert "&lt;script&gt;" in rendered


def test_label_overrides_and_optional_sections_are_applied_without_mutating_source() -> None:
    parsed = parsed_fixture()
    output = generate_report(
        parsed,
        ReportSettings(
            label_overrides={"score": "Assessment score"},
            role_overrides={"score": "numeric"},
            units={"score": "points"},
            valid_ranges={"score": (0.0, 8.0)},
            sections=["overview"],
        ),
    )
    assert "score" in parsed.data.columns
    assert any(column.label == "Assessment score" for column in output.analysis.column_metadata)
    assert output.analysis.quality["outside_user_ranges"]["Assessment score"] == 1
    overview = next(section for section in output.report.sections if section.id == "overview")
    assert any("Assessment score (points)" in row for row in overview.tables[0].rows)
    section_ids = [section.id for section in output.report.sections]
    assert "overview" in section_ids
    assert "quality" not in section_ids and "figures" not in section_ids
    assert {"executive", "limitations", "technical"}.issubset(section_ids)


def test_duplicate_display_labels_are_rejected() -> None:
    parsed = parsed_fixture()
    with pytest.raises(Exception, match="unique"):
        generate_report(parsed, ReportSettings(label_overrides={"score": "post"}))
