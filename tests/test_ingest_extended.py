from __future__ import annotations

from io import BytesIO
import zipfile

import numpy as np
import pandas as pd
import pytest
from openpyxl import Workbook

from analysis import AnalysisError
from inspector import ingest
from inspector.ingest import infer_roles, inspect_xlsx, parse_csv, parse_xlsx


def workbook_bytes(*, multiple: bool = False, formula: bool = False) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Data"
    sheet.append(["participant_id", "score", "group"])
    sheet.append([1, "=40+2" if formula else 42, "A"])
    sheet.append([2, 55, "B"])
    if multiple:
        extra = workbook.create_sheet("Follow-up")
        extra.append(["id", "post"])
        extra.append([1, 48])
    buffer = BytesIO()
    workbook.save(buffer)
    workbook.close()
    return buffer.getvalue()


@pytest.mark.parametrize(
    ("payload", "separator"),
    [
        (b"\xef\xbb\xbfgroup,score\nA,1\n", None),
        ("group;score\nA;1\n".encode(), ";"),
        ("group\tscore\nA\t1\n".encode(), "\t"),
    ],
)
def test_csv_delimiters_and_bom(payload: bytes, separator: str | None) -> None:
    parsed = parse_csv(payload, separator=separator)
    assert list(parsed.data.columns) == ["group", "score"]
    assert parsed.data.loc[0, "score"] == 1


def test_csv_gb18030_and_decimal_comma() -> None:
    parsed = parse_csv("组别;分数\n甲;1,5\n".encode("gb18030"), encoding="gb18030", separator=";", decimal=",")
    assert parsed.data.loc[0, "分数"] == pytest.approx(1.5)


@pytest.mark.parametrize("payload", [b"a,a\n1,2\n", b"a, \n1,2\n", b"a,b\n1\n"])
def test_csv_rejects_duplicate_blank_and_malformed_headers(payload: bytes) -> None:
    with pytest.raises(AnalysisError):
        parse_csv(payload, separator=",")


def test_csv_preserves_na_category_but_marks_blank_missing() -> None:
    parsed = parse_csv(b"code,value\nNA,1\n,2\n", separator=",")
    assert parsed.data.loc[0, "code"] == "NA"
    assert pd.isna(parsed.data.loc[1, "code"])


def test_csv_limits_are_checked_before_dataframe(monkeypatch: pytest.MonkeyPatch) -> None:
    monkeypatch.setattr(ingest, "MAX_ROWS", 1)
    with pytest.raises(AnalysisError, match="row limit"):
        parse_csv(b"a,b\n1,2\n3,4\n", separator=",")


def test_xlsx_single_and_multiple_sheet_selection() -> None:
    one = workbook_bytes()
    assert inspect_xlsx(one) == ["Data"]
    assert parse_xlsx(one).selected_sheet == "Data"
    many = workbook_bytes(multiple=True)
    assert inspect_xlsx(many) == ["Data", "Follow-up"]
    parsed = parse_xlsx(many, sheet_name="Follow-up")
    assert list(parsed.data.columns) == ["id", "post"]


def test_xlsx_formula_cache_disclosure_and_no_evaluation() -> None:
    parsed = parse_xlsx(workbook_bytes(formula=True))
    assert any("cached" in note for note in parsed.notes)
    assert pd.isna(parsed.data.loc[0, "score"])


def test_xlsx_rejects_empty_and_zip_member_limit(monkeypatch: pytest.MonkeyPatch) -> None:
    empty = Workbook()
    buffer = BytesIO(); empty.save(buffer); empty.close()
    with pytest.raises(AnalysisError, match="No visible worksheet"):
        inspect_xlsx(buffer.getvalue())
    monkeypatch.setattr(ingest, "MAX_ZIP_MEMBERS", 1)
    with pytest.raises(AnalysisError, match="member"):
        inspect_xlsx(workbook_bytes())


def test_role_inference_is_conservative_and_excludes_numeric_id_boolean() -> None:
    data = pd.DataFrame({
        "participant_id": [1, 2, 3, 4, 5],
        "measurement": [101.1, 105.2, 109.3, 115.4, 118.5],
        "consent": [True, False, True, True, False],
        "scale": [1, 2, 3, 4, 5],
        "date": ["2026-01-01", "2026-01-02", "2026-01-03", "2026-01-04", "2026-01-05"],
        "free_text": ["long response one", "two", "three", "four", "five"],
    })
    ids = {name: f"c{index}" for index, name in enumerate(data.columns)}
    roles = {column.label: column for column in infer_roles(data, ids)}
    assert roles["participant_id"].role == "identifier"
    assert roles["measurement"].role == "numeric"
    assert roles["consent"].role == "categorical"
    assert roles["scale"].role == "ordinal" and roles["scale"].role_uncertain
    assert roles["date"].role == "date"


def test_role_override_is_explicit() -> None:
    data = pd.DataFrame({"score": [1, 2, 3, 4, 5]})
    role = infer_roles(data, {"score": "c1"}, {"score": "numeric"})[0]
    assert role.role == "numeric"
    assert "override" in role.notes[0].lower()
