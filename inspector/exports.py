"""DOCX and standalone printable HTML exporters for one ReportDocument."""

from __future__ import annotations

import base64
import hashlib
import html
from io import BytesIO
from pathlib import Path
import uuid
import zipfile

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import etree

from .models import Language, ReportDocument, TableBlock


LATIN_FONT = "Arial"
CJK_FONT = "Noto Sans CJK SC"
FONT_PATH = Path(__file__).resolve().parents[1] / "assets" / "fonts" / "NotoSansCJKsc-Regular.otf"
EXPORT_TEXT: dict[Language, dict[str, str]] = {
    "en": {"table": "Table", "continued": "continued", "figure": "Figure", "note": "Note.", "apa": "APA 7–aligned academic reporting", "print": "Open this file in a browser and choose Print → Save as PDF. Client fonts and pagination may vary."},
    "zh-CN": {"table": "表", "continued": "续", "figure": "图", "note": "注：", "apa": "遵循 APA 第 7 版的学术报告", "print": "请在浏览器中打开此文件，并选择“打印 → 另存为 PDF”。分页和客户端字体可能不同。"},
    "es": {"table": "Tabla", "continued": "continuación", "figure": "Figura", "note": "Nota.", "apa": "Informe académico alineado con APA 7", "print": "Abra este archivo en un navegador y elija Imprimir → Guardar como PDF. Las fuentes y la paginación pueden variar."},
    "fi": {"table": "Taulukko", "continued": "jatkuu", "figure": "Kuvio", "note": "Huom.", "apa": "APA 7 -periaatteiden mukainen akateeminen raportointi", "print": "Avaa tiedosto selaimessa ja valitse Tulosta → Tallenna PDF-muodossa. Fontit ja sivutus voivat vaihdella."},
}


def _x(language: Language, key: str) -> str:
    return EXPORT_TEXT[language][key]


def _set_run_font(run, *, size: float | None = None, bold: bool | None = None, italic: bool | None = None) -> None:
    run.font.name = LATIN_FONT
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), CJK_FONT)


def _repeat_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    marker = OxmlElement("w:tblHeader")
    marker.set(qn("w:val"), "true")
    tr_pr.append(marker)


def _set_cell_style(cell, *, fill: str | None = None, padding: int = 90) -> None:
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    tc_pr = cell._tc.get_or_add_tcPr()
    if fill:
        shading = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        if shading.getparent() is None:
            tc_pr.append(shading)
    margins = tc_pr.find(qn("w:tcMar")) or OxmlElement("w:tcMar")
    if margins.getparent() is None:
        tc_pr.append(margins)
    for edge in ("top", "left", "bottom", "right"):
        node = margins.find(qn(f"w:{edge}")) or OxmlElement(f"w:{edge}")
        node.set(qn("w:w"), str(padding))
        node.set(qn("w:type"), "dxa")
        if node.getparent() is None:
            margins.append(node)


def _set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders")) or OxmlElement("w:tblBorders")
    if borders.getparent() is None:
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = borders.find(qn(f"w:{edge}")) or OxmlElement(f"w:{edge}")
        node.set(qn("w:val"), "single")
        node.set(qn("w:sz"), "4")
        node.set(qn("w:color"), "D9D9D9")
        if node.getparent() is None:
            borders.append(node)


def _add_table(document: Document, table: TableBlock, number: int, language: Language) -> None:
    chunks = [table.rows[index : index + 35] for index in range(0, len(table.rows), 35)] or [[]]
    for part, rows in enumerate(chunks, 1):
        caption = document.add_paragraph()
        caption.paragraph_format.keep_with_next = True
        if len(table.columns) >= 7 or table.id == "quality_missing":
            caption.paragraph_format.page_break_before = True
        run = caption.add_run(f"{_x(language, 'table')} {number}" + (f" ({_x(language, 'continued')} {part})" if part > 1 else ""))
        _set_run_font(run, bold=True)
        title = caption.add_run(f"\n{table.title}")
        _set_run_font(title, italic=True)
        grid = document.add_table(rows=1, cols=len(table.columns))
        grid.style = "Table Grid"
        _set_table_borders(grid)
        for cell, value in zip(grid.rows[0].cells, table.columns):
            _set_cell_style(cell, fill="1F4E5F")
            run = cell.paragraphs[0].add_run(value)
            _set_run_font(run, bold=True, size=9)
            run.font.color.rgb = RGBColor(255, 255, 255)
            cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        _repeat_header(grid.rows[0])
        for row_index, source in enumerate(rows):
            cells = grid.add_row().cells
            for cell_index, (cell, value) in enumerate(zip(cells, source)):
                _set_cell_style(cell, fill="EFF6F6" if row_index % 2 else "FFFFFF")
                run = cell.paragraphs[0].add_run(value)
                _set_run_font(run, size=9)
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT if cell_index == 0 else WD_ALIGN_PARAGRAPH.CENTER
        if table.note and part == len(chunks):
            note = document.add_paragraph()
            run = note.add_run(f"{_x(language, 'note')} {table.note}")
            _set_run_font(run, size=9)


def _embed_cjk_font(docx_bytes: bytes) -> bytes:
    """Embed the bundled OFL font so CJK report text remains portable."""

    if not FONT_PATH.exists():
        return docx_bytes
    font_bytes = FONT_PATH.read_bytes()
    font_guid = uuid.UUID(bytes=hashlib.md5(font_bytes).digest())
    font_key = "{" + str(font_guid).upper() + "}"
    xor_key = font_guid.bytes[::-1]
    obfuscated = bytearray(font_bytes)
    for index in range(min(32, len(obfuscated))):
        obfuscated[index] ^= xor_key[index % 16]

    source = BytesIO(docx_bytes)
    target = BytesIO()
    with zipfile.ZipFile(source) as archive, zipfile.ZipFile(target, "w", zipfile.ZIP_DEFLATED) as output:
        font_table = etree.fromstring(archive.read("word/fontTable.xml"))
        word_ns = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
        package_rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
        content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
        relationship_path = "word/_rels/fontTable.xml.rels"
        if relationship_path in archive.namelist():
            relationships = etree.fromstring(archive.read(relationship_path))
        else:
            relationships = etree.Element(f"{{{package_rel_ns}}}Relationships", nsmap={None: package_rel_ns})
        used_ids = {node.get("Id") for node in relationships}
        relationship_id = "rIdCjkFont"
        suffix = 1
        while relationship_id in used_ids:
            suffix += 1
            relationship_id = f"rIdCjkFont{suffix}"
        relation = etree.SubElement(relationships, f"{{{package_rel_ns}}}Relationship")
        relation.set("Id", relationship_id)
        relation.set("Type", "http://schemas.openxmlformats.org/officeDocument/2006/relationships/font")
        relation.set("Target", "fonts/NotoSansCJKsc-Regular.odttf")

        font = etree.SubElement(font_table, f"{{{word_ns}}}font")
        font.set(f"{{{word_ns}}}name", CJK_FONT)
        family = etree.SubElement(font, f"{{{word_ns}}}family")
        family.set(f"{{{word_ns}}}val", "swiss")
        embed = etree.SubElement(font, f"{{{word_ns}}}embedRegular")
        embed.set(f"{{{rel_ns}}}id", relationship_id)
        embed.set(f"{{{word_ns}}}fontKey", font_key)

        content_types = etree.fromstring(archive.read("[Content_Types].xml"))
        if not any(node.get("Extension") == "odttf" for node in content_types):
            content_type = etree.SubElement(content_types, f"{{{content_ns}}}Default")
            content_type.set("Extension", "odttf")
            content_type.set("ContentType", "application/vnd.openxmlformats-officedocument.obfuscatedFont")

        replacements = {
            "word/fontTable.xml": etree.tostring(font_table, xml_declaration=True, encoding="UTF-8", standalone=True),
            relationship_path: etree.tostring(relationships, xml_declaration=True, encoding="UTF-8", standalone=True),
            "[Content_Types].xml": etree.tostring(content_types, xml_declaration=True, encoding="UTF-8", standalone=True),
        }
        for item in archive.infolist():
            if item.filename not in replacements and item.filename != relationship_path:
                output.writestr(item, archive.read(item.filename))
        for name, value in replacements.items():
            output.writestr(name, value)
        output.writestr("word/fonts/NotoSansCJKsc-Regular.odttf", bytes(obfuscated))
    return target.getvalue()


def export_docx(report: ReportDocument) -> bytes:
    document = Document()
    section = document.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    document.settings.element.append(OxmlElement("w:embedTrueTypeFonts"))
    normal = document.styles["Normal"]
    normal.font.name = LATIN_FONT
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    for style_name in ("Title", "Heading 1", "Heading 2"):
        style = document.styles[style_name]
        style.font.name = LATIN_FONT
        style.font.color.rgb = RGBColor(0, 0, 0)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), CJK_FONT)
        p_pr = style._element.get_or_add_pPr()
        border = p_pr.find(qn("w:pBdr"))
        if border is not None:
            p_pr.remove(border)
    title = document.add_paragraph(style="Title")
    _set_run_font(title.add_run(report.title), size=22, bold=True)
    subtitle = document.add_paragraph()
    _set_run_font(subtitle.add_run(report.subtitle), size=11, italic=True)
    status = document.add_paragraph()
    _set_run_font(status.add_run(f"{report.status_note} · {_x(report.language, 'apa')}"), size=9)
    table_number = 0
    figure_number = 0
    for section_model in report.sections:
        heading = document.add_heading(section_model.heading, level=1)
        if section_model.id == "interpretation" and figure_number:
            heading.paragraph_format.page_break_before = True
        for paragraph in section_model.paragraphs:
            document.add_paragraph(paragraph)
        for table in section_model.tables:
            table_number += 1
            _add_table(document, table, table_number, report.language)
        for figure in section_model.figures:
            figure_number += 1
            caption = document.add_paragraph()
            caption.paragraph_format.page_break_before = True
            caption.paragraph_format.keep_with_next = True
            _set_run_font(caption.add_run(f"{_x(report.language, 'figure')} {figure_number}\n"), bold=True)
            _set_run_font(caption.add_run(figure.title), italic=True)
            picture_paragraph = document.add_paragraph()
            picture_paragraph.paragraph_format.keep_together = True
            picture_paragraph.add_run().add_picture(BytesIO(figure.png), width=Inches(5.9))
            note = document.add_paragraph()
            _set_run_font(note.add_run(f"{_x(report.language, 'note')} {figure.caption}"), size=9)
    output = BytesIO()
    document.save(output)
    serialized = output.getvalue()
    return _embed_cjk_font(serialized) if report.language == "zh-CN" else serialized


def _html_table(table: TableBlock, number: int, language: Language) -> str:
    headers = "".join(f"<th scope='col'>{html.escape(value)}</th>" for value in table.columns)
    rows = "".join("<tr>" + "".join(f"<td>{html.escape(value)}</td>" for value in row) + "</tr>" for row in table.rows)
    note = f"<p class='note'><em>{html.escape(_x(language, 'note'))}</em> {html.escape(table.note)}</p>" if table.note else ""
    return f"<figure class='table'><figcaption><strong>{html.escape(_x(language, 'table'))} {number}</strong><br><em>{html.escape(table.title)}</em></figcaption><div class='table-wrap'><table><thead><tr>{headers}</tr></thead><tbody>{rows}</tbody></table></div>{note}</figure>"


def export_html(report: ReportDocument) -> bytes:
    table_number = 0
    figure_number = 0
    parts: list[str] = []
    for section in report.sections:
        body = [f"<h2>{html.escape(section.heading)}</h2>"]
        body.extend(f"<p>{html.escape(paragraph)}</p>" for paragraph in section.paragraphs)
        for table in section.tables:
            table_number += 1
            body.append(_html_table(table, table_number, report.language))
        for figure in section.figures:
            figure_number += 1
            encoded = base64.b64encode(figure.png).decode("ascii")
            body.append(f"<figure><figcaption><strong>{html.escape(_x(report.language, 'figure'))} {figure_number}</strong><br><em>{html.escape(figure.title)}</em></figcaption><img src='data:image/png;base64,{encoded}' alt='{html.escape(figure.alt_text, quote=True)}'><p class='note'><em>{html.escape(_x(report.language, 'note'))}</em> {html.escape(figure.caption)}</p></figure>")
        parts.append("<section>" + "".join(body) + "</section>")
    document = f"""<!doctype html>
<html lang="{html.escape(report.language)}"><head><meta charset="utf-8"><meta name="viewport" content="width=device-width,initial-scale=1">
<title>{html.escape(report.title)}</title>
<style>
:root{{--ink:#172323;--muted:#526464;--accent:#167d7f;--line:#cfdada}}*{{box-sizing:border-box}}body{{margin:0;background:#f5f7f6;color:var(--ink);font-family:"Noto Sans CJK SC","Noto Sans SC","Microsoft YaHei",Arial,sans-serif;line-height:1.58}}main{{max-width:900px;margin:32px auto;padding:48px 58px;background:white;box-shadow:0 6px 25px #10202014}}h1{{font-size:2rem;margin:0}}h2{{margin-top:2.2rem;border-bottom:2px solid var(--accent);padding-bottom:.25rem}}.subtitle,.status,.note{{color:var(--muted)}}figure{{margin:1.5rem 0;break-inside:avoid}}figcaption{{margin-bottom:.6rem}}img{{display:block;max-width:100%;height:auto}}.table-wrap{{overflow-x:auto}}table{{width:100%;border-collapse:collapse;font-size:.92rem}}th{{text-align:left;border-bottom:2px solid var(--ink);padding:.45rem}}td{{border-bottom:1px solid var(--line);padding:.4rem}}.print-help{{border-left:4px solid var(--accent);padding:.6rem 1rem;background:#edf7f7}}@page{{margin:18mm}}@media print{{body{{background:white}}main{{margin:0;max-width:none;padding:0;box-shadow:none}}.print-help{{display:none}}h2{{break-after:avoid}}table{{font-size:9pt}}}}
</style></head><body><main><header><h1>{html.escape(report.title)}</h1><p class="subtitle">{html.escape(report.subtitle)}</p><p class="status">{html.escape(report.status_note)} · {html.escape(_x(report.language, 'apa'))}</p><p class="print-help">{html.escape(_x(report.language, 'print'))}</p></header>{''.join(parts)}</main></body></html>"""
    return document.encode("utf-8")
