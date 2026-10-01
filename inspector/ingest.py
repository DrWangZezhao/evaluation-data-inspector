"""Bounded CSV/XLSX ingestion and conservative variable-role inference."""

from __future__ import annotations

import csv
import re
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from io import BytesIO, StringIO
from pathlib import Path
from typing import Literal

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from pandas.api.types import is_bool_dtype, is_datetime64_any_dtype, is_numeric_dtype

from analysis import AnalysisError, MAX_COLUMNS, MAX_ROWS, MAX_UPLOAD_BYTES, validate_dataframe
from .models import ColumnMeta, VariableRole


MAX_CELLS = 2_000_000
MAX_ZIP_MEMBERS = 2_000
MAX_ZIP_UNCOMPRESSED = 100 * 1024 * 1024
ENCODINGS = {"UTF-8": "utf-8-sig", "GB18030": "gb18030"}
SEPARATORS = {"Automatic": None, "Comma": ",", "Semicolon": ";", "Tab": "\t"}


@dataclass
class ParsedData:
    data: pd.DataFrame
    sheet_names: list[str] = field(default_factory=list)
    selected_sheet: str | None = None
    notes: list[str] = field(default_factory=list)
    column_ids: dict[str, str] = field(default_factory=dict)


def _check_headers(headers: list[object]) -> list[str]:
    labels = ["" if value is None else str(value).strip() for value in headers]
    if not labels or all(not label for label in labels):
        raise AnalysisError("The first row must contain column names.")
    if any(not label for label in labels):
        raise AnalysisError("Every column must have a non-blank header.")
    folded = [label.casefold() for label in labels]
    duplicates = sorted({label for label in folded if folded.count(label) > 1})
    if duplicates:
        raise AnalysisError("Column names must be unique; duplicate headers were found.")
    if len(labels) > MAX_COLUMNS:
        raise AnalysisError(f"The dataset exceeds the {MAX_COLUMNS}-column limit.")
    return labels


def _detect_separator(text: str) -> str:
    lines = [line for line in text.splitlines()[:12] if line.strip()]
    if not lines:
        raise AnalysisError("The CSV is empty.")
    counts = {separator: [line.count(separator) for line in lines] for separator in (",", ";", "\t")}
    viable = [sep for sep, values in counts.items() if min(values) > 0 and len(set(values)) == 1]
    if len(viable) == 1:
        return viable[0]
    if len(viable) > 1:
        best = max(viable, key=lambda sep: counts[sep][0])
        if sum(counts[sep][0] == counts[best][0] for sep in viable) == 1:
            return best
    raise AnalysisError("The CSV separator is ambiguous. Choose comma, semicolon, or tab in Advanced settings.")


def parse_csv(
    content: bytes,
    *,
    encoding: str = "utf-8-sig",
    separator: str | None = None,
    decimal: Literal[".", ","] = ".",
) -> ParsedData:
    if not content:
        raise AnalysisError("The uploaded CSV is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise AnalysisError("The upload exceeds the 10 MiB limit.")
    if b"\x00" in content[:4096]:
        raise AnalysisError("The upload appears to be a binary file, not a CSV.")
    try:
        text = content.decode(encoding)
    except (UnicodeDecodeError, LookupError) as exc:
        raise AnalysisError("The CSV does not match the selected encoding.") from exc
    delimiter = separator or _detect_separator(text)
    try:
        reader = csv.reader(StringIO(text), delimiter=delimiter, strict=True)
        header = next(reader)
        labels = _check_headers(header)
        row_count = 0
        for row in reader:
            row_count += 1
            if len(row) != len(labels):
                raise AnalysisError(f"Malformed CSV row {row_count + 1}: expected {len(labels)} fields.")
            if row_count > MAX_ROWS:
                raise AnalysisError(f"The dataset exceeds the {MAX_ROWS:,}-row limit.")
            if row_count * len(labels) > MAX_CELLS:
                raise AnalysisError(f"The dataset exceeds the {MAX_CELLS:,}-cell limit.")
    except (csv.Error, StopIteration) as exc:
        raise AnalysisError("The CSV header or row structure is malformed.") from exc
    if row_count == 0:
        raise AnalysisError("The CSV must contain at least one data row.")
    try:
        data = pd.read_csv(
            StringIO(text), delimiter=delimiter, decimal=decimal, keep_default_na=False,
            na_values=[""], on_bad_lines="error",
        )
    except (pd.errors.ParserError, pd.errors.EmptyDataError, ValueError) as exc:
        raise AnalysisError("The CSV could not be parsed with the selected options.") from exc
    data.columns = labels
    validate_dataframe(data)
    notes = [f"CSV encoding: {encoding}; separator: {repr(delimiter)}; decimal mark: {decimal}."]
    return ParsedData(data=data, notes=notes, column_ids={label: f"c{i:03d}" for i, label in enumerate(labels, 1)})


def _inspect_zip(content: bytes) -> None:
    try:
        with zipfile.ZipFile(BytesIO(content)) as archive:
            members = archive.infolist()
            if len(members) > MAX_ZIP_MEMBERS:
                raise AnalysisError(f"The workbook exceeds the {MAX_ZIP_MEMBERS:,}-member archive limit.")
            if sum(member.file_size for member in members) > MAX_ZIP_UNCOMPRESSED:
                raise AnalysisError("The workbook expands beyond the 100 MiB safety limit.")
            names = {member.filename.lower() for member in members}
            if "[content_types].xml" not in names:
                raise AnalysisError("The file is not a valid XLSX workbook.")
            if any(name.endswith("vbaproject.bin") for name in names):
                raise AnalysisError("Macro-enabled workbooks are not supported. Save a clean .xlsx copy.")
    except zipfile.BadZipFile as exc:
        raise AnalysisError("The workbook is corrupt, encrypted, or not a supported XLSX file.") from exc


def inspect_xlsx(content: bytes) -> list[str]:
    if not content:
        raise AnalysisError("The uploaded workbook is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise AnalysisError("The upload exceeds the 10 MiB limit.")
    _inspect_zip(content)
    try:
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    except Exception as exc:
        raise AnalysisError("The workbook is corrupt, encrypted, or unsupported.") from exc
    usable: list[str] = []
    try:
        for sheet in workbook.worksheets:
            if sheet.sheet_state != "visible":
                continue
            first = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
            second = next(sheet.iter_rows(min_row=2, max_row=2, values_only=True), ())
            if first and any(value not in (None, "") for value in first) and any(value not in (None, "") for value in second):
                usable.append(sheet.title)
    finally:
        workbook.close()
    if not usable:
        raise AnalysisError("No visible worksheet contains a header row and data row.")
    return usable


def parse_xlsx(content: bytes, *, sheet_name: str | None = None) -> ParsedData:
    usable = inspect_xlsx(content)
    chosen = sheet_name or usable[0]
    if chosen not in usable:
        raise AnalysisError("Choose one of the usable visible worksheets.")
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True, keep_links=False)
    try:
        sheet = workbook[chosen]
        rows = sheet.iter_rows(values_only=True)
        raw_header = next(rows, ())
        while raw_header and raw_header[-1] in (None, ""):
            raw_header = raw_header[:-1]
        labels = _check_headers(list(raw_header))
        records: list[list[object]] = []
        trailing_empty = 0
        for index, row in enumerate(rows, 2):
            values = list(row[: len(labels)])
            if len(values) < len(labels):
                values.extend([None] * (len(labels) - len(values)))
            if all(value in (None, "") for value in values):
                trailing_empty += 1
                continue
            if trailing_empty:
                records.extend([[None] * len(labels) for _ in range(trailing_empty)])
                trailing_empty = 0
            records.append(values)
            if len(records) > MAX_ROWS:
                raise AnalysisError(f"The worksheet exceeds the {MAX_ROWS:,}-row limit.")
            if len(records) * len(labels) > MAX_CELLS:
                raise AnalysisError(f"The worksheet exceeds the {MAX_CELLS:,}-cell limit.")
        if not records:
            raise AnalysisError("The selected worksheet contains no data rows.")
        data = pd.DataFrame.from_records(records, columns=labels)
        validate_dataframe(data)
    finally:
        workbook.close()
    notes = [
        f"Worksheet: {chosen}. Trailing empty padding was ignored; internal blank records were preserved.",
        "Formula cells use cached workbook results, which may be stale or absent; formulas were not evaluated.",
        "External workbook links were not loaded.",
    ]
    return ParsedData(data=data, sheet_names=usable, selected_sheet=chosen, notes=notes, column_ids={label: f"c{i:03d}" for i, label in enumerate(labels, 1)})


def parse_upload(
    content: bytes,
    filename: str,
    *,
    sheet_name: str | None = None,
    encoding: str = "utf-8-sig",
    separator: str | None = None,
    decimal: Literal[".", ","] = ".",
) -> ParsedData:
    suffix = Path(filename).suffix.lower()
    if suffix == ".csv":
        return parse_csv(content, encoding=encoding, separator=separator, decimal=decimal)
    if suffix == ".xlsx":
        return parse_xlsx(content, sheet_name=sheet_name)
    raise AnalysisError("Only .csv and .xlsx files are supported.")


_ID_NAME = re.compile(r"(^|[_\s-])(id|identifier|participant.?id|record.?id|case.?id)($|[_\s-])", re.I)
_ISO_DATE = re.compile(r"^\d{4}[-/]\d{1,2}[-/]\d{1,2}(?:[ T].*)?$")


def infer_roles(data: pd.DataFrame, column_ids: dict[str, str], overrides: dict[str, VariableRole] | None = None) -> list[ColumnMeta]:
    overrides = overrides or {}
    result: list[ColumnMeta] = []
    for label in data.columns:
        series = data[label]
        nonmissing = series.dropna()
        unique = int(nonmissing.nunique(dropna=True))
        missing = int(series.isna().sum())
        nonfinite = 0
        uncertain = False
        notes: list[str] = []
        role: VariableRole
        if label in overrides:
            role = overrides[label]
            notes.append("User override applied.")
        elif is_bool_dtype(series):
            role = "categorical"
        elif is_datetime64_any_dtype(series) or any(isinstance(value, (date, datetime)) for value in nonmissing.head(20)):
            role = "date"
        elif is_numeric_dtype(series):
            numeric = pd.to_numeric(series, errors="coerce").astype(float)
            nonfinite = int(np.isinf(numeric).sum())
            if _ID_NAME.search(str(label)) and unique >= max(2, int(len(nonmissing) * 0.8)):
                role = "identifier"
            elif unique <= 10 and len(nonmissing) >= 5 and np.allclose(numeric.dropna() % 1, 0):
                role = "ordinal"
                uncertain = True
                notes.append("Low-cardinality integer scale; confirm order before ordered analysis.")
            else:
                role = "numeric"
        else:
            strings = nonmissing.astype(str).str.strip()
            iso_dates = len(strings) >= 2 and float(strings.str.match(_ISO_DATE).mean()) >= 0.9
            uniqueness = unique / max(1, len(nonmissing))
            if iso_dates:
                role = "date"
                uncertain = True
                notes.append("Unambiguous date-like strings; confirm role if needed.")
            elif _ID_NAME.search(str(label)) and uniqueness >= 0.8:
                role = "identifier"
            elif unique <= 20 or uniqueness <= 0.2:
                role = "categorical"
            else:
                role = "text"
        result.append(ColumnMeta(
            id=column_ids[str(label)], label=str(label), role=role, role_uncertain=uncertain,
            dtype=str(series.dtype), nonmissing=int(len(nonmissing)), unique_nonmissing=unique,
            missing=missing, nonfinite=nonfinite, notes=notes,
        ))
    return result
